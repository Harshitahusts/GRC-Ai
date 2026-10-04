"""The AI checks that an evidence file is about the DPDPA obligation it's linked to."""

import io
import json
from types import SimpleNamespace

import pytest
from docx import Document
from helpers import create, csrf, post

from grc_agent import evidence_check
from grc_agent.config import Settings
from grc_agent.web import db

BREACH_POLICY = (
    "Personal data breach response procedure. On becoming aware of a breach we inform "
    "each affected Data Principal and send the Data Protection Board an intimation without "
    "delay, then a detailed report within 72 hours."
)


class FakeAI:
    """Answers like a provider would; records what it was asked."""

    def __init__(self, verdict="relevant", reason="It is a breach procedure.", raw=None):
        self.calls = []
        self.answer = raw or json.dumps(
            {"verdict": verdict, "reason": reason, "missing": ["Board intimation template"]}
        )
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self.create))

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(content=[SimpleNamespace(text=self.answer)], stop_reason="end_turn")


@pytest.fixture
def ai(app):
    fake = FakeAI()
    app.state.ai = Settings(provider="groq", model="openai/gpt-oss-120b")
    app.state.ai_client = fake
    return fake


def _upload(authed, eid, name="breach.txt", data=None, obligation="OBL-007"):
    data = BREACH_POLICY.encode() if data is None else data
    return authed.post(
        f"/engagements/{eid}/evidence",
        data={
            "csrf": csrf(authed, "/"),
            "title": "Breach procedure",
            "category": "Policy or procedure",
            "obligation_id": obligation,
        },
        files={"file": (name, data, "text/plain")},
    )


def _file(app):
    with db.connect(app.state.db_path) as conn:
        return conn.execute("SELECT * FROM evidence_files ORDER BY id DESC").fetchone()


def test_upload_is_checked_automatically_and_labelled(authed, app, ai):
    eid = create(authed)
    page = _upload(authed, eid).text
    assert "The AI is checking" in page
    f = _file(app)
    assert f["ai_check"] == "relevant" and "breach procedure" in f["ai_check_reason"]
    assert "Groq" in f["ai_checked_by"]
    prompt = ai.calls[0]["messages"][0]["content"]
    assert "72 hours" in prompt and "<document>" in prompt
    page = authed.get(f"/engagements/{eid}/evidence").text
    assert "AI check: On topic" in page and "Board intimation template" in page


def test_an_off_topic_file_stops_counting_until_a_person_overrules(authed, app, ai):
    ai.answer = json.dumps(
        {"verdict": "not_relevant", "reason": "This is a CV, not a policy.", "missing": []}
    )
    eid = create(authed)
    oid = "OBL-007"
    _upload(authed, eid, obligation=oid)
    fid = _file(app)["id"]
    page = authed.get(f"/engagements/{eid}/evidence").text
    assert "Not about this obligation" in page and "Doesn't count as evidence" in page

    # Marking the control implemented needs real evidence (or a written explanation).
    r = post(authed, f"/engagements/{eid}/controls/{oid}", {"status": "implemented"})
    assert "attach evidence or describe" in r.text

    post(authed, f"/engagements/{eid}/evidence/{fid}/overrule")
    assert _file(app)["check_overruled_by"] == "harshit"
    r = post(authed, f"/engagements/{eid}/controls/{oid}", {"status": "implemented"})
    assert "Implemented" in r.text and "attach evidence" not in r.text
    with db.connect(app.state.db_path) as conn:
        actions = [r[0] for r in conn.execute("SELECT action FROM audit_log")]
    assert "evidence_checked" in actions and "evidence_check_overruled" in actions


def test_check_again_by_hand(authed, app, ai):
    eid = create(authed)
    _upload(authed, eid)
    fid = _file(app)["id"]
    ai.answer = json.dumps(
        {"verdict": "partly_relevant", "reason": "No Board report step.", "missing": []}
    )
    page = post(authed, f"/engagements/{eid}/evidence/{fid}/check").text
    assert "AI check of Breach procedure: Partly on topic" in page
    assert _file(app)["ai_check"] == "partly_relevant"


def test_no_check_without_an_ai_provider_or_obligation(authed, app):
    eid = create(authed)
    page = _upload(authed, eid).text  # no key set up: nothing runs
    assert "The AI is checking" not in page and _file(app)["ai_check"] == ""
    fid = _file(app)["id"]
    page = post(authed, f"/engagements/{eid}/evidence/{fid}/check").text
    assert "Couldn&#39;t check" in page


def test_unusable_answer_and_unreadable_files_are_reported(authed, app, ai):
    eid = create(authed)
    ai.answer = "I think it's fine!"
    _upload(authed, eid)
    assert _file(app)["ai_check"] == ""  # background: left unchecked, no crash
    fid = _file(app)["id"]
    assert "couldn&#39;t be read" in post(authed, f"/engagements/{eid}/evidence/{fid}/check").text

    png = authed.post(
        f"/engagements/{eid}/evidence",
        data={"csrf": csrf(authed, "/"), "obligation_id": "OBL-007"},
        files={"file": ("shot.png", b"\x89PNG fake", "image/png")},
    )
    assert png.status_code == 200
    fid = _file(app)["id"]
    page = post(authed, f"/engagements/{eid}/evidence/{fid}/check").text
    assert "Images and spreadsheets can&#39;t be read" in page
    # Recorded on the file without an AI call, so the row says why it wasn't checked.
    assert _file(app)["ai_check"] == "unreadable" and "no AI call" in _file(app)["ai_checked_by"]
    assert "Couldn&#39;t read the text" in authed.get(f"/engagements/{eid}/evidence").text


def test_text_extraction_from_word_and_the_injection_guard():
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Consent notice")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Purpose"
    doc.save(buf)
    text = evidence_check.extract_text(buf.getvalue(), ".docx")
    assert "Consent notice" in text and "Purpose" in text
    with pytest.raises(evidence_check.Unreadable):
        evidence_check.extract_text(b"%PDF-1.4 no text", ".pdf")
    assert (
        "never" in evidence_check.SYSTEM_PROMPT and "instructions" in evidence_check.SYSTEM_PROMPT
    )


def test_tiny_files_are_judged_without_calling_the_ai():
    fake = FakeAI()
    result = evidence_check.check(
        fake, Settings(provider="groq", model="m"), None, "t", "Other", "hi"
    )
    assert result.verdict == "too_little_content" and not fake.calls


def test_read_the_text_the_app_extracted(authed, app):
    eid = create(authed)
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Breach response: tell the Board within 72 hours <b>now</b>.")
    doc.save(buf)
    authed.post(
        f"/engagements/{eid}/evidence",
        data={"csrf": csrf(authed, "/"), "title": "Breach SOP", "obligation_id": "OBL-005"},
        files={"file": ("sop.docx", buf.getvalue(), "application/octet-stream")},
    )
    fid = _file(app)["id"]
    listing = authed.get(f"/engagements/{eid}/evidence").text
    assert f"/engagements/{eid}/evidence/{fid}/text" in listing
    page = authed.get(f"/engagements/{eid}/evidence/{fid}/text").text
    assert "tell the Board within 72 hours" in page
    assert "&lt;b&gt;now&lt;/b&gt;" in page  # shown as text, never as HTML
    with db.connect(app.state.db_path) as conn:
        actions = [r[0] for r in conn.execute("SELECT action FROM audit_log")]
    assert "evidence_text_viewed" in actions
    # Another engagement can't read it.
    other = create(authed)
    assert authed.get(f"/engagements/{other}/evidence/{fid}/text").status_code == 404
